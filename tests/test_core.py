import os
from pathlib import Path
import subprocess
import sys

import numpy as np
import pytest
from scipy.sparse import coo_matrix
from docx import Document

from Corpus import Corpus, prepare_data
from GloVe import GloVe, cosine
from lower_tokenizer import lower_tokenizer, sp_remover
from color_docx import generate_docx
from extend_docx import partial_docx, whole_docx

ROOT = Path(__file__).resolve().parents[1]


def pairs(count=8.0):
    return coo_matrix(([count], ([0], [1])), shape=(2, 2))


def fitted(seed=0, epochs=2):
    model = GloVe(dim=2, seed=seed).fit_vectors(pairs(), epochs=epochs)
    model.add_dictionary({"a": 0, "b": 1})
    return model


def test_seed_zero_is_respected():
    a, b = fitted(0), fitted(0)
    assert a.seed == 0
    np.testing.assert_array_equal(a.word_vectors, b.word_vectors)
    np.testing.assert_array_equal(a.word_biases, b.word_biases)
    assert not np.array_equal(a.word_vectors, fitted(1).word_vectors)


def test_one_step_matches_independent_gradient():
    # Expected simultaneous gradients use the same PRE-update vector pair.
    old = (np.random.RandomState(0).rand(2, 2) - 0.5) / 2
    residual = (8 / 100)**0.5 * (np.dot(old[0], old[1]) - np.log(8))
    expected = old.copy()
    expected[0] -= 0.02 * residual * old[1]
    expected[1] -= 0.02 * residual * old[0]
    model = fitted(epochs=1)
    np.testing.assert_allclose(model.word_vectors, expected, rtol=0, atol=1e-15)
    np.testing.assert_allclose(model.word_biases, [-0.02 * residual] * 2)
    np.testing.assert_allclose(model.biases_sum_gradients, [1 + residual**2] * 2)
    np.testing.assert_allclose(model.vectors_sum_gradients[0], 1 + (residual * old[1])**2)
    np.testing.assert_allclose(model.vectors_sum_gradients[1], 1 + (residual * old[0])**2)


def test_training_reduces_toy_objective():
    def loss(model):
        residual = np.dot(*model.word_vectors) + sum(model.word_biases) - np.log(8)
        return 0.5 * (8 / 100)**0.5 * residual**2
    assert loss(fitted(epochs=40)) < loss(fitted(epochs=1))


@pytest.mark.parametrize("count", [0, -1, np.nan, np.inf])
def test_invalid_counts(count):
    with pytest.raises(ValueError):
        GloVe().fit_vectors(pairs(count))


@pytest.mark.parametrize("matrix", [np.eye(2), coo_matrix((2, 2)),
    coo_matrix(([1.0], ([0], [1])), shape=(2, 3)),
    coo_matrix(([1.0], ([0], [0])), shape=(2, 2))])
def test_invalid_matrices(matrix):
    with pytest.raises(ValueError):
        GloVe().fit_vectors(matrix)


@pytest.mark.parametrize("kwargs", [{"dim": 0}, {"dim": 1.5}, {"seed": -1},
    {"seed": 2**32}, {"seed": 0.5}, {"learning_rate": 0}, {"alpha": np.nan},
    {"max_count": -1}, {"max_loss": 0}])
def test_invalid_parameters(kwargs):
    with pytest.raises(ValueError):
        GloVe(**kwargs)


@pytest.mark.parametrize("epochs", [0, -1, 0.5, True])
def test_invalid_epochs(epochs):
    with pytest.raises(ValueError):
        GloVe().fit_vectors(pairs(), epochs=epochs)


def test_zero_vector_and_shape_handling():
    assert cosine([0, 0], [1, 0]) == 0
    assert cosine([1, 0], [1, 0]) == 1
    assert cosine([1, 0], [-1, 0]) == -1
    with pytest.raises(ValueError):
        cosine([1], [1, 2])


def test_tied_neighbors_exclude_query_itself():
    model = fitted()
    model.word_vectors[:] = [1, 0]
    assert model.get_most_similar("b", 10) == [("a", 1.0)]
    assert model.get_most_similar("a", 0) == []


def test_missing_words_and_untrained_model():
    model = fitted()
    assert model.get_most_similar("missing") == []
    assert model.check_similarity("a", "missing") == -1
    with pytest.raises(KeyError):
        model.check_similarity("a", "missing", ignore_missing=False)
    with pytest.raises(ValueError):
        GloVe().get_most_similar("a")


def test_dictionary_validation_and_copy():
    model = fitted()
    for dictionary in ({"a": 0}, {"a": 0, "b": 0}, {"a": 2, "b": 3}):
        with pytest.raises(ValueError):
            model.add_dictionary(dictionary)
    dictionary = {"a": 0, "b": 1}
    model.add_dictionary(dictionary)
    dictionary["extra"] = 2
    assert "extra" not in model.dictionary


def test_corpus_instances_do_not_share_dictionary():
    a, b = Corpus(), Corpus()
    a.fit_matrix([["a", "b"]])
    assert b.dictionary == {}
    dictionary = {"a": 0}
    c = Corpus(dictionary)
    c.fit_matrix([["a", "b"]])
    assert dictionary == {"a": 0}


def test_cooccurrence_counts_are_hand_checkable():
    corpus = Corpus()
    corpus.fit_matrix([["a", "b", "c"]], window_size=2)
    np.testing.assert_array_equal(corpus.matrix.toarray(), [[0, 1, 0.5], [0, 0, 1], [0, 0, 0]])
    with pytest.raises(ValueError):
        corpus.fit_matrix([["a", "b"]], window_size=0)


def test_offline_tokenization_and_imports(tmp_path):
    code = '''import nltk
nltk.data.path = []
def forbidden(*args, **kwargs):
    raise RuntimeError("unexpected download")
nltk.download = forbidden
import Corpus, lower_tokenizer, train_model
assert lower_tokenizer.lower_tokenizer("HELLO world") == ["hello", "world"]
assert Corpus.tokenizer("Hello world") == ["Hello", "world"]
'''
    environment = dict(os.environ, PYTHONPATH=str(ROOT))
    result = subprocess.run([sys.executable, "-c", code], cwd=tmp_path, env=environment,
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    assert result.stdout == ""
    assert sp_remover(["the", "cat"], stop_words={"the"}) == ["cat"]
    assert prepare_data("A B", tokenizer=lower_tokenizer) == [["a", "b"]]


def test_model_and_corpus_round_trip(tmp_path):
    model = fitted()
    model.save_file(tmp_path / "model.pkl")
    restored = GloVe.load_file(tmp_path / "model.pkl")
    np.testing.assert_array_equal(restored.word_vectors, model.word_vectors)
    assert restored.get_most_similar("a") == model.get_most_similar("a")
    corpus = Corpus()
    corpus.fit_matrix([["a", "b"]])
    corpus.save_file(tmp_path / "corpus.pkl")
    restored_corpus = Corpus.load_file(tmp_path / "corpus.pkl")
    assert restored_corpus.dictionary == corpus.dictionary
    np.testing.assert_array_equal(restored_corpus.matrix.toarray(), corpus.matrix.toarray())


def test_document_helpers_handle_empty_and_preserve_input(tmp_path, monkeypatch):
    monkeypatch.chdir(tmp_path)
    generate_docx([], [])
    assert Document("colored_docx.docx").paragraphs == []
    words = ["hello", "world"]
    generate_docx(words, [0.95, 0.5])
    assert words == ["hello", "world"]
    assert Document("colored_docx.docx").paragraphs[0].text == "Hello world"
    doc = Document()
    partial_docx([], [], doc)
    assert len(doc.paragraphs) == 1
    partial_docx(words, [0.9, 0.5], doc)
    assert words == ["hello", "world"]


def test_document_shape_validation():
    with pytest.raises(ValueError):
        generate_docx(["a"], [])
    with pytest.raises(ValueError):
        partial_docx(["a"], [], Document())
    with pytest.raises(ValueError):
        whole_docx([["a"]], [])
    with pytest.raises(ValueError):
        generate_docx([""], [0.5])


def test_cli_runs_without_existing_models(tmp_path):
    result = subprocess.run([sys.executable, str(ROOT / "train_model.py"), "--epochs", "2",
                             "--output", str(tmp_path / "output")], cwd=tmp_path,
                            capture_output=True, text=True, timeout=30)
    assert result.returncode == 0, result.stderr
    assert (tmp_path / "output/model.pkl").is_file()
    assert (tmp_path / "output/corpus.pkl").is_file()
    assert "seed=0" in result.stdout


def test_cli_missing_corpus_fails_cleanly(tmp_path):
    result = subprocess.run([sys.executable, str(ROOT / "train_model.py"), "--corpus",
                             str(tmp_path / "missing.txt")], capture_output=True, text=True, timeout=30)
    assert result.returncode == 2
    assert "Traceback" not in result.stderr
